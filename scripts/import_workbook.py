from __future__ import annotations

"""Import a personal-finance workbook without altering the original file."""

import argparse
from collections import defaultdict
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
import shutil
import sys
from typing import Any

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

import app.database as database  # noqa: E402


PERSIAN_MONTHS = {
    "فروردین": 1,
    "اردیبهشت": 2,
    "خرداد": 3,
    "تیر": 4,
    "مرداد": 5,
    "شهریور": 6,
    "مهر": 7,
    "آبان": 8,
    "آذر": 9,
    "دی": 10,
    "بهمن": 11,
    "اسفند": 12,
}
PAID_STATUSES = {"پرداخت شده", "تسویه شده"}


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def amount_value(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        cleaned = value.replace(",", "").replace("٫", ".").strip()
        if not cleaned or cleaned.startswith("="):
            return None
        value = cleaned
    try:
        number = int(float(value))
    except (TypeError, ValueError):
        return None
    return number if number >= 0 else None


def iso_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def workbook_digest(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def archive_all_rows(db, import_id: int, workbook) -> int:
    total = 0
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(), start=1):
            values = {
                cell.column_letter: json_value(cell.value)
                for cell in row
                if cell.value is not None and cell.value != ""
            }
            if not values:
                continue
            db.execute(
                """INSERT INTO imported_rows(import_id, sheet_name, source_row, record_type, data_json)
                   VALUES (?, ?, ?, 'raw', ?)""",
                (import_id, sheet.title, row_number, json.dumps(values, ensure_ascii=False, default=str)),
            )
            total += 1
    return total


def find_header_row(sheet, expected_header: str) -> int | None:
    for row in sheet.iter_rows():
        for cell in row:
            if cell_text(cell.value) == expected_header:
                return cell.row
    return None


def import_budgets(db, import_id: int, values_workbook) -> int:
    imported = 0
    for sheet in values_workbook.worksheets:
        if not sheet.title.startswith("بودجه "):
            continue
        try:
            fiscal_year = int(sheet.title.split()[-1])
        except ValueError:
            continue
        header_row = find_header_row(sheet, "سرفصل بودجه")
        if not header_row:
            continue
        headers = {
            cell.column: PERSIAN_MONTHS.get(cell_text(cell.value))
            for cell in sheet[header_row]
        }
        for row_number in range(header_row + 1, sheet.max_row + 1):
            category = cell_text(sheet.cell(row_number, 1).value)
            if not category:
                continue
            for column, month in headers.items():
                if not month:
                    continue
                amount = amount_value(sheet.cell(row_number, column).value)
                if amount is None:
                    continue
                source_key = f"budget:{sheet.title}:{row_number}:{column}"
                db.execute(
                    """INSERT OR IGNORE INTO budget_items
                       (fiscal_year, category_name, jalali_month, amount, import_id, source_key)
                       VALUES (?, ?, ?, ?, ?, ?)""",
                    (fiscal_year, category, month, amount, import_id, source_key),
                )
                imported += 1
    return imported


def import_commitments(db, import_id: int, values_workbook) -> tuple[int, int, int]:
    if "تعهدات مالی" not in values_workbook.sheetnames:
        return 0, 0, 0
    sheet = values_workbook["تعهدات مالی"]
    header_row = find_header_row(sheet, "نوع")
    if not header_row:
        return 0, 0, 0

    groups: dict[tuple[str, int | None, str], list[tuple[int, tuple[Any, ...]]]] = defaultdict(list)
    for row_number in range(header_row + 1, sheet.max_row + 1):
        title = cell_text(sheet.cell(row_number, 2).value)
        total = amount_value(sheet.cell(row_number, 3).value)
        original_date = cell_text(sheet.cell(row_number, 4).value)
        due_date = iso_date(sheet.cell(row_number, 5).value)
        amount = amount_value(sheet.cell(row_number, 8).value)
        status = cell_text(sheet.cell(row_number, 7).value)
        payment_month = cell_text(sheet.cell(row_number, 10).value)
        if not title or not due_date or not amount or amount <= 0:
            continue
        groups[(title, total, original_date)].append(
            (row_number, (due_date, amount, status, payment_month))
        )

    commitment_count = installment_count = payment_count = 0
    for group_index, ((title, total, original_date), rows) in enumerate(groups.items(), start=1):
        commitment_key = f"commitment:financial-obligations:{group_index}:{title}:{total}:{original_date}"
        db.execute(
            """INSERT OR IGNORE INTO commitments(title, kind, total_amount, source_key, unique_code)
               VALUES (?, ?, ?, ?, ?)""",
            (title, "تعهد مالی", total, commitment_key, database.allocate_unique_code(db)),
        )
        commitment_id = db.execute(
            "SELECT id FROM commitments WHERE source_key = ?", (commitment_key,)
        ).fetchone()["id"]
        commitment_count += 1
        for row_number, (due_date, amount, status, payment_month) in rows:
            source_key = f"installment:تعهدات مالی:{row_number}"
            note = ""
            if payment_month:
                note += f"ماه پرداخت ثبت‌شده در اکسل: {payment_month}"
            db.execute(
                """INSERT OR IGNORE INTO installments
                   (commitment_id, due_date, amount, note, source_key)
                   VALUES (?, ?, ?, ?, ?)""",
                (commitment_id, due_date, amount, note, source_key),
            )
            installment = db.execute(
                "SELECT id FROM installments WHERE source_key = ?", (source_key,)
            ).fetchone()
            if not installment:
                continue
            installment_count += 1
            if status in PAID_STATUSES:
                payment_key = f"payment:تعهدات مالی:{row_number}"
                db.execute(
                    """INSERT OR IGNORE INTO payments
                       (installment_id, amount, paid_on, note, source_key)
                       VALUES (?, ?, ?, ?, ?)""",
                    (
                        installment["id"],
                        amount,
                        due_date,
                        "تاریخ پرداخت در فایل موجود نبود، سررسید به‌عنوان تاریخ ثبت شد.",
                        payment_key,
                    ),
                )
                payment_count += 1
    return commitment_count, installment_count, payment_count


def backup_database() -> Path | None:
    if not database.DATABASE_PATH.exists():
        return None
    return database.create_database_backup(database.DATA_DIR / "backups")


def run_import(workbook_path: Path, create_backup: bool = True) -> dict[str, Any]:
    workbook_path = workbook_path.expanduser().resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    digest = workbook_digest(workbook_path)
    backup = backup_database() if create_backup else None
    database.initialize_database()

    with database.connection() as db:
        existing = db.execute(
            "SELECT id, status FROM import_runs WHERE workbook_hash = ?", (digest,)
        ).fetchone()
        if existing and existing["status"] == "completed":
            return {"already_imported": True, "import_id": existing["id"], "backup": str(backup) if backup else None}

        formula_workbook = load_workbook(workbook_path, data_only=False, read_only=False)
        values_workbook = load_workbook(workbook_path, data_only=True, read_only=False)
        cursor = db.execute(
            """INSERT INTO import_runs(source_name, source_path, workbook_hash, status)
               VALUES (?, ?, ?, 'running')""",
            (workbook_path.name, str(workbook_path), digest),
        )
        import_id = cursor.lastrowid
        total_rows = archive_all_rows(db, import_id, formula_workbook)
        commitment_count, installment_count, payment_count = import_commitments(db, import_id, values_workbook)
        budget_count = import_budgets(db, import_id, values_workbook)
        db.execute(
            """UPDATE import_runs
               SET status = 'completed', total_rows = ?, mapped_commitments = ?, mapped_installments = ?,
                   mapped_payments = ?, mapped_budget_items = ?, completed_at = CURRENT_TIMESTAMP
               WHERE id = ?""",
            (total_rows, commitment_count, installment_count, payment_count, budget_count, import_id),
        )

    return {
        "already_imported": False,
        "import_id": import_id,
        "total_rows": total_rows,
        "commitments": commitment_count,
        "installments": installment_count,
        "payments": payment_count,
        "budget_items": budget_count,
        "backup": str(backup) if backup else None,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Import an Excel workbook into Financial Assistant.")
    parser.add_argument("workbook", type=Path, help="Path to the Excel workbook")
    args = parser.parse_args()
    result = run_import(args.workbook)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
