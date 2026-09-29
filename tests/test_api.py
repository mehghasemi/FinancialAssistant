import asyncio
from contextlib import closing
from concurrent.futures import ThreadPoolExecutor
import logging
import sqlite3
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import database, main

logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("httpx2").setLevel(logging.WARNING)


class ApiTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.directory = Path(temp.name)
        for name, value in (("DATA_DIR", self.directory), ("DATABASE_PATH", self.directory / "test.db")):
            mock = patch.object(database, name, value)
            mock.start()
            self.addCleanup(mock.stop)
        self.client = self.enterContext(TestClient(main.app, raise_server_exceptions=False))

    def create(self, **changes):
        payload = dict(title="وام نمونه", kind="وام", installment_amount=1000,
                       installment_count=3, first_due_date="1405/07/01", interval_months=2)
        payload.update(changes)
        response = self.client.post("/api/commitments", json=payload)
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["id"]

    def rows(self, identifier):
        return self.client.get("/api/installments", params={"commitment_id": identifier}).json()

    def test_commitment_summary_and_details_keep_same_title_loans_separate(self):
        first = self.create()
        second = self.create()
        rows = self.rows(first)
        for row, amount in ((rows[0], 1000), (rows[1], 400)):
            response = self.client.post('/api/payments', json={
                'installment_id': row['id'], 'amount': amount, 'paid_on': '1405/07/01'})
            self.assertEqual(response.status_code, 201)
        summaries = {item['id']: item for item in self.client.get('/api/commitments').json()}
        self.assertEqual(summaries[first]['paid_installment_count'], 1)
        self.assertEqual(summaries[first]['paid_amount'], 1400)
        self.assertEqual(summaries[first]['planned_amount'], 3000)
        self.assertEqual(summaries[second]['paid_installment_count'], 0)
        self.assertEqual(summaries[second]['paid_amount'], 0)
        self.assertEqual(len(self.rows(first)), 3)
        self.assertTrue(all(row['commitment_id'] == first for row in self.rows(first)))
        self.assertTrue(all(row['paid_amount'] == 0 for row in self.rows(second)))

    def edit(self, identifier, **changes):
        return self.client.patch(f"/api/commitments/{identifier}", json={"title": "وام جدید", "kind": "وام", **changes})

    def test_summary_counts_overdue_partial_payments_and_variable_amounts(self):
        identifier = self.create(first_due_date='1400/01/01', repayment_amount=3100)
        rows = self.rows(identifier)
        for row, amount in ((rows[0], 1000), (rows[1], 400)):
            response = self.client.post('/api/payments', json={
                'installment_id': row['id'], 'amount': amount, 'paid_on': '1400/01/01'})
            self.assertEqual(response.status_code, 201)
        item = next(item for item in self.client.get('/api/commitments').json() if item['id'] == identifier)
        self.assertEqual(item['overdue_installment_count'], 2)
        self.assertEqual(item['paid_installment_count'], 1)
        self.assertEqual(item['installment_amount_variants'], 2)
        self.assertEqual(item['planned_amount'] - item['paid_amount'], 1700)
        self.assertLess(item['first_due_date'], item['last_due_date'])

    def test_installment_edit_adjusts_repayment_and_audits_before_after(self):
        identifier = self.create(repayment_amount=3000)
        rows = self.rows(identifier)
        payload = {'due_date': rows[0]['due_date'], 'amount': 1300, 'note': 'اصلاح'}
        response = self.client.patch(f"/api/installments/{rows[0]['id']}", json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        after = self.rows(identifier)
        self.assertEqual(after[1:], rows[1:])
        item = self.client.get('/api/commitments').json()[0]
        self.assertEqual(item['repayment_amount'], 3300)
        logs = self.client.get('/api/audit-logs', params={'resource_type': 'installment', 'resource_id': rows[0]['id']}).json()['items']
        self.assertEqual(logs[0]['details']['before']['amount'], 1000)
        self.assertEqual(logs[0]['details']['after']['amount'], 1300)
        payload['amount'] = 800
        self.assertEqual(self.client.patch(f"/api/installments/{rows[0]['id']}", json=payload).status_code, 200)
        self.assertEqual(self.client.get('/api/commitments').json()[0]['repayment_amount'], 2800)
        payload['note'] = 'فقط یادداشت'
        self.assertEqual(self.client.patch(f"/api/installments/{rows[0]['id']}", json=payload).status_code, 200)
        self.assertEqual(self.client.get('/api/commitments').json()[0]['repayment_amount'], 2800)

    def test_audit_pagination_and_settings_details(self):
        self.create()
        payload = self.client.get('/api/settings').json()
        payload['backup_enabled'] = False
        self.assertEqual(self.client.put('/api/settings', json=payload).status_code, 200)
        first = self.client.get('/api/audit-logs?limit=1').json()
        self.assertEqual(first['items'][0]['resource_type'], 'settings')
        self.assertEqual(first['items'][0]['details']['after']['backup_enabled'], 'false')
        self.assertIsNotNone(first['next_cursor'])
        second = self.client.get('/api/audit-logs', params={'before_id': first['next_cursor']}).json()
        self.assertTrue(all(item['id'] < first['next_cursor'] for item in second['items']))

    def test_paid_schedule_protected_and_metadata_edit_preserves_history(self):
        identifier = self.create()
        rows = self.rows(identifier)
        payment = self.client.post("/api/payments", json={
            "installment_id": rows[0]["id"], "amount": 400, "paid_on": "1405/07/01"})
        self.assertEqual(payment.status_code, 201)
        self.assertEqual(self.client.post("/api/payments", json={
            "installment_id": rows[-1]["id"], "amount": 100, "paid_on": "1405/07/01"}).status_code, 201)
        self.assertEqual(self.edit(identifier, installment_count=2).status_code, 422)
        self.assertEqual(self.edit(identifier, repayment_amount=600).status_code, 422)
        self.assertEqual(self.edit(identifier, first_due_date="1405/08/01").status_code, 200)
        self.assertEqual(self.edit(identifier, total_amount=5000, installment_count=3,
                                   repayment_amount=3000, interval_months=2, first_due_date="1405/07/01").status_code, 200)
        updated = self.rows(identifier)
        self.assertEqual([r["id"] for r in updated], [r["id"] for r in rows])
        self.assertEqual(updated[0]["paid_amount"], 400)
        detail = self.client.get(f'/api/installments/{rows[0]["id"]}/details').json()
        self.assertEqual(detail["payments"][0]["id"], payment.json()["id"])
        self.assertTrue(any(item["resource_type"] == "commitment" for item in detail["history"]))

    def test_all_editable_fields_propagate_and_paid_installments_keep_identity(self):
        identifier = self.create(repayment_amount=3000)
        other = self.create(title="وام نمونه", repayment_amount=3000)
        before = self.rows(identifier)
        self.client.post("/api/payments", json={"installment_id": before[0]["id"], "amount": 1000, "paid_on": "1405/07/01"})
        response = self.edit(identifier, kind="بدهی", total_amount=5000, repayment_amount=6500,
                             installment_count=4, installment_amount=1500, interval_months=1, first_due_date="1405/08/01")
        self.assertEqual(response.status_code, 200, response.text)
        after = self.rows(identifier)
        self.assertEqual([r["amount"] for r in after], [1500, 1500, 1500, 2000])
        self.assertEqual([r["id"] for r in after[:3]], [r["id"] for r in before])
        self.assertTrue(all(r["title"] == "وام جدید" and r["kind"] == "بدهی" and r["total_amount"] == 5000 for r in after))
        self.assertEqual(after[0]["paid_amount"], 1000)
        self.assertEqual(after[-1]["due_date"], "۱۴۰۵/۱۱/۰۱")
        self.assertEqual(self.rows(other)[0]["title"], "وام نمونه")
        details = self.client.get(f'/api/installments/{before[0]["id"]}/details').json()
        self.assertEqual(details["payments"][0]["paid_on"], "۱۴۰۵/۰۷/۰۱")

    def test_preview_and_create_agree_on_repayment_and_last_installment(self):
        payload = dict(title="وام نمونه", kind="وام", total_amount=2000, repayment_amount=3100,
                       installment_count=3, installment_amount=1000, first_due_date="1405/07/01")
        preview = self.client.post("/api/commitments/preview", json=payload)
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual(preview.json()["planned_total"], 3100)
        self.assertEqual([r["amount"] for r in preview.json()["installments"]], [1000, 1000, 1100])
        identifier = self.create(**payload)
        self.assertEqual([r["amount"] for r in self.rows(identifier)], [1000, 1000, 1100])
        listed = self.client.get("/api/commitments").json()[0]
        self.assertEqual(listed["repayment_amount"], 3100)
        self.assertEqual(listed["installment_amount"], 1000)
        self.assertTrue(listed["first_due_date"])
        payload["installment_amount"] = 2000
        self.assertEqual(self.client.post("/api/commitments/preview", json=payload).status_code, 422)

    def test_date_only_edit_preserves_unequal_amounts_and_payment(self):
        identifier = self.create(repayment_amount=3100)
        before = self.rows(identifier)
        self.client.post("/api/payments", json={"installment_id": before[0]["id"], "amount": 400, "paid_on": "1405/07/01"})
        self.assertEqual(self.edit(identifier, first_due_date="1405/08/01").status_code, 200)
        after = self.rows(identifier)
        self.assertEqual([r["amount"] for r in after], [r["amount"] for r in before])
        self.assertEqual(after[0]["paid_amount"], 400)

    def test_schedule_changes_preserve_ids_and_exact_total(self):
        identifier = self.create()
        before = self.rows(identifier)
        self.assertEqual(self.edit(identifier, installment_count=4, repayment_amount=4003,
                                   interval_months=1, first_due_date="1405/08/01").status_code, 200)
        after = self.rows(identifier)
        self.assertEqual([r["id"] for r in after[:3]], [r["id"] for r in before])
        self.assertEqual([r["amount"] for r in after], [1000, 1000, 1000, 1003])
        self.assertEqual(after[0]["due_date"], "۱۴۰۵/۰۸/۰۱")
        self.assertEqual(self.edit(identifier, installment_count=2).status_code, 200)
        self.assertEqual(sum(r["amount"] for r in self.rows(identifier)), 4003)

    def test_invalid_schedule_rolls_back(self):
        identifier = self.create()
        before = self.rows(identifier)
        self.assertEqual(self.edit(identifier, repayment_amount=2).status_code, 422)
        self.assertEqual(self.rows(identifier), before)

    def test_date_only_and_interval_only_updates(self):
        identifier = self.create()
        self.assertEqual(self.edit(identifier, first_due_date="1405/08/01").status_code, 200)
        self.assertEqual(self.rows(identifier)[0]["due_date"], "۱۴۰۵/۰۸/۰۱")
        self.assertEqual(self.edit(identifier, interval_months=1).status_code, 200)
        self.assertEqual(self.rows(identifier)[1]["due_date"], "۱۴۰۵/۰۹/۰۱")

    def test_http_validation_not_found_and_account_errors(self):
        self.assertEqual(self.client.post("/api/commitments", json={}).status_code, 422)
        self.assertEqual(self.edit(999).status_code, 404)
        identifier = self.create()
        installment = self.rows(identifier)[0]["id"]
        payload = {"installment_id": installment, "amount": 100, "paid_on": "1405/07/01", "account_id": 999}
        self.assertEqual(self.client.post("/api/payments", json=payload).status_code, 422)
        payload.update(account_id=1, amount=1001)
        self.assertEqual(self.client.post("/api/payments", json=payload).status_code, 422)
        self.assertEqual(self.rows(identifier)[0]["paid_amount"], 0)

    def test_lifespan_creates_startup_backup_and_respects_setting(self):
        backups = list((self.directory / "backups").glob("*.db"))
        self.assertTrue(backups)
        self.client.put("/api/settings", json={"backup_enabled": False, "backup_directory": str(self.directory / "backups")})
        main.automatic_backup()
        self.assertEqual(list((self.directory / "backups").glob("*.db")), backups)

    def test_clear_data_requires_confirmation_and_keeps_backup(self):
        identifier = self.create()
        for payload in ({}, {"confirmation": "no"}):
            self.assertEqual(self.client.post("/api/settings/clear-data", json=payload).status_code, 422)
        self.assertTrue(self.rows(identifier))
        settings = self.client.get("/api/settings").json()
        with database.connection(write=True) as db:
            db.execute("INSERT INTO import_runs(source_name, source_path, workbook_hash, status) VALUES ('test', '', 'hash', 'completed')")
            db.execute("INSERT INTO imported_rows(import_id, sheet_name, source_row, data_json) VALUES (1, 'Sheet4', 1, '{}')")
            db.execute("INSERT INTO budget_items(fiscal_year, category_name, jalali_month, amount, import_id) VALUES (1405, 'test', 1, 100, 1)")
            db.execute("INSERT INTO transactions(transaction_type, amount, occurred_on) VALUES ('income', 100, '2026-09-29')")
        response = self.client.post("/api/settings/clear-data", json={"confirmation": "DELETE_ALL_FINANCIAL_DATA"})
        self.assertEqual(response.status_code, 200, response.text)
        with closing(sqlite3.connect(response.json()["backup_path"])) as backup:
            self.assertEqual(backup.execute("SELECT count(*) FROM commitments").fetchone()[0], 1)
        with database.connection() as db:
            for table in ("payments", "installments", "commitments", "transactions", "budget_items", "imported_rows", "import_runs"):
                self.assertEqual(db.execute(f"SELECT count(*) FROM {table}").fetchone()[0], 0, table)
            self.assertTrue(db.execute("SELECT 1 FROM audit_logs WHERE resource_type = 'archived_commitment'").fetchone())
            self.assertTrue(db.execute("SELECT 1 FROM audit_logs WHERE resource_type = 'system'").fetchone())
            self.assertTrue(db.execute("SELECT 1 FROM data_migrations").fetchone())
            self.assertEqual(db.execute("PRAGMA foreign_key_check").fetchall(), [])
        self.assertEqual(self.client.get("/api/settings").json(), settings)
        self.create()

    def test_clear_data_failure_preserves_records(self):
        identifier = self.create()
        payload = {"confirmation": "DELETE_ALL_FINANCIAL_DATA"}
        with patch.object(database, "create_database_backup", side_effect=OSError("backup failed")):
            with self.assertLogs(main.logger, level="ERROR"):
                self.assertEqual(self.client.post("/api/settings/clear-data", json=payload).status_code, 500)
        self.assertEqual(len(self.rows(identifier)), 3)
        with database.connection(write=True) as db:
            db.execute("CREATE TRIGGER block_clear BEFORE DELETE ON commitments BEGIN SELECT RAISE(ABORT, 'blocked'); END")
        with self.assertLogs(main.logger, level="ERROR"):
            self.assertEqual(self.client.post("/api/settings/clear-data", json=payload).status_code, 500)
        self.assertEqual(len(self.rows(identifier)), 3)

    def test_concurrent_payments_cannot_exceed_balance(self):
        identifier = self.create()
        installment = self.rows(identifier)[0]["id"]
        payload = {"installment_id": installment, "amount": 700, "paid_on": "1405/07/01"}
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: self.client.post("/api/payments", json=payload).status_code, range(2)))
        self.assertEqual(sorted(responses), [201, 422])
        self.assertEqual(self.rows(identifier)[0]["paid_amount"], 700)

    def test_internal_error_is_sanitized(self):
        with patch.object(main.commitments.finance, "create_commitment", side_effect=RuntimeError("private database detail")):
            with self.assertLogs(main.logger, level="ERROR"):
                response = self.client.post("/api/commitments", json={
                    "title": "وام نمونه", "kind": "وام", "installment_amount": 1000,
                    "installment_count": 1, "first_due_date": "1405/07/01"})
        self.assertEqual(response.status_code, 500)
        self.assertNotIn("private database detail", response.text)
        self.assertEqual(response.headers["X-Error-Id"], response.json()["error_id"])

    def test_downloaded_template_is_empty_and_can_be_imported_after_filling(self):
        url = '/static/templates/installments-template.xlsx'
        self.assertIn(url, self.client.get('/').text)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.addCleanup(workbook.close)
        sheet = workbook['Sheet4']
        self.assertEqual([cell.value for cell in sheet[1]],
                         ['عنوان', 'نوع', 'مبلغ کل', 'تاریخ سررسید', 'وضعیت', 'مبلغ'])
        self.assertFalse(any(cell.value is not None for row in sheet.iter_rows(min_row=2) for cell in row))
        self.assertIn('راهنما و نمونه', workbook.sheetnames)
        empty = self.client.post('/api/imports/sheet4', files={'file': ('template.xlsx', response.content)})
        self.assertEqual(empty.status_code, 422)
        self.assertEqual(self.client.get('/api/commitments').json(), [])
        sheet.append(['تعهد قالب', 'وام', 20000000, '1405/07/01', 'پرداخت شده', 11000000])
        sheet.append(['تعهد قالب', 'وام', 20000000, '1405/08/01', 'پرداخت نشده', 11000000])
        output = BytesIO()
        workbook.save(output)
        result = self.client.post('/api/imports/sheet4', files={'file': ('template.xlsx', output.getvalue())})
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['installments'], 2)
        self.assertEqual(result.json()['payments'], 1)
        records = self.client.get('/api/commitments').json()
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]['repayment_amount'], 2200000)
        self.assertEqual(records[0]['paid_amount'], 1100000)

    def upload_rows(self, rows, headers=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Sheet4'
        sheet.append(headers or ['عنوان', 'نوع', 'مبلغ کل', 'تاریخ سررسید', 'وضعیت', 'مبلغ'])
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return self.client.post('/api/imports/sheet4', files={'file': ('test.xlsx', output.getvalue())})

    def test_excel_validation_reports_rows_without_partial_writes(self):
        valid = ['وام', 'وام', 20000, '1405/07/01', 'پرداخت نشده', 10000]
        response = self.upload_rows([valid, ['', 'وام', -10, '1405/13/01', 'نامعلوم', 11]])
        self.assertEqual(response.status_code, 422)
        message = response.json()['detail']
        for expected in ['هیچ اطلاعاتی ثبت نشد', 'ردیف 3', 'عنوان', 'مبلغ کل', 'وضعیت', 'تاریخ سررسید']:
            self.assertIn(expected, message)
        self.assertEqual(self.client.get('/api/commitments').json(), [])

    def test_excel_rejects_invalid_numbers_formulas_duplicates_and_headers(self):
        valid = ['وام', 'وام', 20000, '1405/07/01', 'پرداخت نشده', 10000]
        for amount in ['NaN', 'Infinity', 10.5, 0, -10, 9, 11, True, '=10000', '9999999999999999999999990']:
            with self.subTest(amount=amount):
                self.assertEqual(self.upload_rows([valid[:-1] + [amount]]).status_code, 422)
        response = self.upload_rows([valid, valid])
        self.assertEqual(response.status_code, 422)
        self.assertIn('قسط تکراری', response.json()['detail'])
        response = self.upload_rows([valid], ['عنوان', 'عنوان', 'مبلغ کل', 'تاریخ سررسید', 'وضعیت', 'مبلغ'])
        self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.get('/api/commitments').json(), [])

    def test_excel_existing_collision_rejected_and_identical_retry_allowed(self):
        valid = ['وام', 'وام', 20000, '1405/07/01', 'پرداخت نشده', 10000]
        self.assertEqual(self.upload_rows([valid]).status_code, 200)
        before = self.client.get('/api/commitments').json()
        self.assertEqual(self.upload_rows([valid]).status_code, 200)
        changed = valid.copy()
        changed[3] = '1405/08/01'
        response = self.upload_rows([changed])
        self.assertEqual(response.status_code, 422)
        self.assertIn('ردیف 2', response.json()['detail'])
        self.assertEqual(self.client.get('/api/commitments').json(), before)

    def test_excel_ignores_empty_rows_but_rejects_incomplete_rows(self):
        valid = ['وام', '', 20000, '1405/07/01', 'پرداخت نشده', 10000]
        self.assertEqual(self.upload_rows([[None] * 6, valid]).status_code, 200)
        response = self.upload_rows([[None, None, None, '1405/07/01', None, None]])
        self.assertEqual(response.status_code, 422)
        self.assertIn('ردیف 2', response.json()['detail'])

    def test_excel_upload_remains_available_without_legacy_folder(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet4"
        sheet.append(["عنوان", "نوع", "مبلغ کل", "تاریخ سررسید", "وضعیت", "مبلغ"])
        sheet.append(["تعهد اکسل", "وام", 20000, "1405/08/01", "پرداخت شده", 10000])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        for _ in range(2):
            response = self.client.post("/api/imports/sheet4", files={
                "file": ("sample.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            self.assertEqual(response.status_code, 200, response.text)
        commitments = self.client.get("/api/commitments").json()
        self.assertEqual(len(commitments), 1)
        self.assertEqual(commitments[0]["paid_amount"], 1000)
        self.assertEqual(commitments[0]["installment_count"], 1)
        self.assertIn('id="importSheet4File"', self.client.get("/").text)
        cleared = self.client.post("/api/settings/clear-data", json={"confirmation": "DELETE_ALL_FINANCIAL_DATA"})
        self.assertEqual(cleared.status_code, 200)
        response = self.client.post("/api/imports/sheet4", files={"file": ("sample.xlsx", output.getvalue())})
        self.assertEqual(response.status_code, 200, response.text)
        imported = self.client.get("/api/commitments").json()
        self.assertEqual(len(imported), 1)
        self.assertEqual(imported[0]["paid_amount"], 1000)


class BackupLoopTests(unittest.TestCase):
    def test_periodic_backup_and_clean_stop(self):
        async def exercise():
            stop = asyncio.Event()
            def backup():
                loop.call_soon_threadsafe(stop.set)
            loop = asyncio.get_running_loop()
            with patch.object(main, "BACKUP_INTERVAL_SECONDS", 0.001), patch.object(main, "automatic_backup", side_effect=backup) as action:
                await asyncio.wait_for(main.periodic_backups(stop), timeout=2)
                self.assertEqual(action.call_count, 1)
        asyncio.run(exercise())
